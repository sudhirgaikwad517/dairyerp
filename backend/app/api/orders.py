import io

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy.orm import Session
from pydantic import BaseModel
from typing import List, Optional
from datetime import date, datetime
from backend.app.database import get_db
from backend.app.models.entities import Customer, Delivery, DeliveryRoute, DeliveryTrip, Order, OrderItem, Product, User
from backend.app.services.attachments import attachments_for
from backend.app.core.security import require_admin
from backend.app.services.billing_service import bill_for_delivery_day, get_business_profile
from backend.app.services.daily_orders import dispatch_day, order_deadline, set_day_quantity
from backend.app.services.dispatch_control import (DISPATCH_LOCK, check_manual_dispatch, dispatch_status,
                                                    release_hold, set_hold)
from backend.app.services.order_lines import describe, litres_per_unit, order_lines, order_what

router = APIRouter(prefix="/orders", tags=["Orders"])

# Orders whose milk has been delivered and is billed
BILLED_STATUSES = ("delivered", "invoiced", "paid")


class OrderItemInput(BaseModel):
    product_id: int
    quantity: float


class HoldSchema(BaseModel):
    day: date
    note: str


class CreateOrderSchema(BaseModel):
    customer_id: int
    delivery_date: Optional[str] = None
    items: List[OrderItemInput]
    notes: Optional[str] = None


def _parse_day(value: Optional[str]) -> date:
    try:
        return date.fromisoformat(value) if value else date.today()
    except ValueError:
        raise HTTPException(status_code=400, detail="Date must be YYYY-MM-DD")


def _order_dict(db: Session, o: Order, products, routes, drivers) -> dict:
    items = db.query(OrderItem).filter(OrderItem.order_id == o.id).all()
    delivery = db.query(Delivery).filter(Delivery.order_id == o.id).order_by(Delivery.id.desc()).first()
    trip = db.query(DeliveryTrip).filter(DeliveryTrip.id == delivery.trip_id).first() if delivery else None
    customer = o.customer
    route = routes.get(trip.route_id if trip else (customer.assigned_route_id if customer else None))
    driver = drivers.get(trip.driver_id if trip else (customer.assigned_driver_id if customer else None))

    delivered = o.status in BILLED_STATUSES
    # Same per-product rule as the bill: only what reached the customer is billed
    rows = order_lines(db, [o], products).get(o.id, [])
    delivered_litres = round(sum(r["delivered_litres"] for r in rows), 3) if delivered else None
    bill = bill_for_delivery_day(db, o.customer_id, o.delivery_date) if delivered else None

    items_detail = []
    for r in rows:
        items_detail.append({
            "item_id": r["item"].id,
            "product_id": r["product_id"],
            "product_name": r["product_name"],
            "quantity": r["ordered"],
            "delivered_quantity": r["delivered"] if delivered else None,
            "unit": r["unit"],
            "unit_price": r["rate"],
            "total_amount": r["ordered_amount"],
            "billed_amount": r["amount"] if delivered else None,
        })
    short_items = [r for r in rows if delivered and r["delivered"] < r["ordered"]]
    rates = sorted({i.unit_price for i in items})
    return {
        "id": o.id,
        "order_number": o.order_number,
        "customer_id": o.customer_id,
        "customer_name": customer.business_name if customer else "Customer",
        "customer_code": customer.customer_code if customer else "",
        "route_name": route.route_name if route else None,
        "driver_name": driver.full_name if driver else None,
        "order_date": o.order_date.isoformat() if o.order_date else None,
        "delivery_date": o.delivery_date.isoformat(),
        "status": o.status,
        "total_litres": o.total_litres,
        "delivered_litres": delivered_litres,
        "what": describe((r["ordered"], r["product"]) for r in rows) if o.status != "cancelled" else "",
        "delivered_what": describe((r["delivered"], r["product"]) for r in rows) if delivered else None,
        "is_short": bool(short_items),
        "short_what": describe((round(r["ordered"] - r["delivered"], 3), r["product"]) for r in short_items) if short_items else None,
        "multi_product": len(rows) > 1,
        "delivery_status": delivery.status if delivery else None,
        "shortage_reason": delivery.shortage_reason if delivery and delivery.shortage_reason not in (None, "none") else None,
        "rate": rates[0] if len(rates) == 1 and len(rows) == 1 else None,
        "subtotal": o.subtotal,
        "transport_charge": o.transport_charge,
        "total_amount": o.total_amount,
        # What this day adds to the bill (delivered milk only)
        "billed_amount": round(sum(r["amount"] for r in rows), 2) if delivered else None,
        "bill_number": bill.invoice_number if bill else None,
        "notes": o.notes,
        # Proof of delivery (delivery boy's photo / customer signature)
        "delivery_id": delivery.id if delivery else None,
        "delivered_at": delivery.departure_time.isoformat() if delivery and delivery.departure_time and delivered else None,
        "delivered_to": delivery.delivered_to_person if delivery else None,
        "photo_ids": [a.id for a in attachments_for(db, "delivery", delivery.id)] if delivery else [],
        "has_signature": bool(delivery and delivery.customer_signature_url),
        "items": items_detail,
    }


@router.get("/")
def list_orders(
    customer_id: Optional[int] = None,
    order_date: Optional[str] = None,
    status: Optional[str] = None,
    limit: int = 500,
    db: Session = Depends(get_db)
):
    """Orders by delivery date (order_date param = delivery date), newest first."""
    query = db.query(Order).order_by(Order.delivery_date.desc(), Order.id.desc())
    if customer_id:
        query = query.filter(Order.customer_id == customer_id)
    if order_date:
        query = query.filter(Order.delivery_date == _parse_day(order_date))
    if status:
        query = query.filter(Order.status == status)
    products = {p.id: p for p in db.query(Product).all()}
    routes = {r.id: r for r in db.query(DeliveryRoute).all()}
    drivers = {u.id: u for u in db.query(User).filter(User.role == "driver").all()}
    return [_order_dict(db, o, products, routes, drivers) for o in query.limit(max(1, min(limit, 2000))).all()]


@router.get("/day-summary")
def day_summary(day: Optional[str] = None, db: Session = Depends(get_db)):
    """One delivery day at a glance: who ordered, litres, dispatch progress and who has not ordered."""
    d = _parse_day(day)
    orders = db.query(Order).filter(Order.delivery_date == d).order_by(Order.id.desc()).all()
    latest = {}
    for o in orders:  # one order per customer per day; newest wins for old data
        latest.setdefault(o.customer_id, o)
    active = [o for o in latest.values() if o.status != "cancelled" and (o.total_amount or 0) > 0]
    products = {p.id: p for p in db.query(Product).all()}
    per_product = {}
    for rows in order_lines(db, active, products).values():
        for r in rows:
            p = products.get(r["product_id"])
            t = per_product.setdefault(r["product_id"], {
                "product_id": r["product_id"], "product_name": r["product_name"], "unit": r["unit"],
                # Sizes of one product (1 L, 500 ML) share a group name; the app shows them together
                "group_name": (p.group_name if p and p.group_name else r["product_name"]),
                "variant_name": (p.variant_name or "") if p else "",
                "quantity": 0.0, "litres": 0.0, "amount": 0.0, "customers": 0})
            t["quantity"] = round(t["quantity"] + r["ordered"], 3)
            t["litres"] = round(t["litres"] + r["ordered"] * litres_per_unit(p), 3)
            t["amount"] = round(t["amount"] + r["ordered_amount"], 2)
            t["customers"] += 1
    by_product = sorted(per_product.values(), key=lambda t: ((products[t["product_id"]].sort_order or 0) if t["product_id"] in products else 0, t["product_id"]))
    counts = {}
    for o in active:
        counts[o.status] = counts.get(o.status, 0) + 1
    customers = db.query(Customer).filter(Customer.is_active == True).order_by(Customer.customer_code).all()
    ordered_ids = {o.customer_id for o in active}
    deadline = order_deadline(db, d)
    return {
        "date": d.isoformat(),
        "active_customers": len(customers),
        "customers_ordered": len(ordered_ids),
        "not_ordered": [{"id": c.id, "name": c.business_name} for c in customers if c.id not in ordered_ids],
        "cancelled": sum(1 for o in latest.values() if o.status == "cancelled"),
        "total_litres": round(sum(o.total_litres for o in active), 3),
        "total_amount": round(sum(o.total_amount for o in active), 2),
        "products": by_product,
        "status_counts": counts,
        "waiting_dispatch": counts.get("created", 0) + counts.get("confirmed", 0),
        "order_deadline": deadline.isoformat(),
        "customers_can_order": datetime.now() < deadline,
        "dispatch": dispatch_status(db, d),
    }


@router.post("/")
def create_order(payload: CreateOrderSchema, db: Session = Depends(get_db)):
    """Set a customer's order for a day (one order per customer per day; placing again changes it)."""
    cust = db.query(Customer).filter(Customer.id == payload.customer_id).first()
    if not cust:
        raise HTTPException(status_code=404, detail="Customer not found")
    deliv_date = _parse_day(payload.delivery_date)
    order = set_day_quantity(db, cust, deliv_date, [(i.product_id, i.quantity) for i in payload.items], placed_by="admin", is_admin=True)

    credit_warning = None
    # Credit limit 0 means no limit has been set for this customer
    if cust.credit_limit and cust.credit_limit > 0 and (cust.current_outstanding or 0) + order.total_amount > cust.credit_limit:
        credit_warning = (f"Notice: over credit limit (Outstanding: ₹{cust.current_outstanding:,.2f}, "
                          f"Limit: ₹{cust.credit_limit:,.2f})")
    return {
        "status": "success",
        "order_id": order.id,
        "order_number": order.order_number,
        "order_status": order.status,
        "total_amount": order.total_amount,
        "total_litres": order.total_litres,
        "what": order_what(db, order) if order.status != "cancelled" else "",
        "credit_warning": credit_warning
    }


@router.post("/dispatch")
def dispatch_orders(day: Optional[str] = None, force: bool = False, db: Session = Depends(get_db)):
    """Put the day's confirmed orders onto route delivery trips so delivery boys can see and deliver them."""
    d = _parse_day(day)
    check_manual_dispatch(db, d)
    deadline = order_deadline(db, d)
    if datetime.now() < deadline and not force:
        raise HTTPException(
            status_code=400,
            detail=f"Customers can still order for {d:%d %b} until {deadline:%d %b %I:%M %p}. "
                   f"Dispatching now locks their orders.")
    with DISPATCH_LOCK:
        return dispatch_day(db, d)


@router.get("/dispatch-status")
def get_dispatch_status(day: Optional[str] = None, db: Session = Depends(get_db)):
    return dispatch_status(db, _parse_day(day))


@router.put("/dispatch-hold")
def hold_dispatch(payload: HoldSchema, admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    """Stop the automatic dispatch for a day; at dispatch time the waiting customers get the note."""
    return set_hold(db, payload.day, payload.note, admin)


@router.delete("/dispatch-hold")
def remove_dispatch_hold(day: str, admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    """Remove the hold; if the dispatch time has passed, the orders go to the delivery boys now."""
    return release_hold(db, _parse_day(day), admin)


@router.get("/export")
def export_day(day: Optional[str] = None, db: Session = Depends(get_db)):
    """Excel sheet of one delivery day's orders: per customer, per route, and who has not ordered."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    d = _parse_day(day)
    summary = day_summary(d.isoformat(), db)
    products = {p.id: p for p in db.query(Product).all()}
    routes = {r.id: r for r in db.query(DeliveryRoute).all()}
    drivers = {u.id: u for u in db.query(User).filter(User.role == "driver").all()}
    latest = {}
    for o in db.query(Order).filter(Order.delivery_date == d).order_by(Order.id.desc()).all():
        latest.setdefault(o.customer_id, o)
    rows = [_order_dict(db, o, products, routes, drivers) for o in latest.values()
            if o.status != "cancelled" and (o.total_amount or 0) > 0]
    cancelled = [_order_dict(db, o, products, routes, drivers) for o in latest.values() if o.status == "cancelled"]
    seq = {c.id: (c.route_stop_sequence or 0) for c in db.query(Customer).all()}
    rows.sort(key=lambda r: (r["route_name"] or "~", seq.get(r["customer_id"], 0), r["customer_name"]))
    customers = {c.id: c for c in db.query(Customer).all()}

    profile = get_business_profile(db)
    deadline = datetime.fromisoformat(summary["order_deadline"])
    disp = summary["dispatch"]
    if summary["customers_can_order"]:
        state = f"Not final - customers can still order until {deadline:%d %b %Y %I:%M %p}"
    else:
        state = f"Final - ordering closed at {deadline:%d %b %Y %I:%M %p}"
    if disp["on_hold"]:
        dispatch_line = f"Dispatch ON HOLD: {disp['hold_note']}"
    else:
        dispatch_line = f"Automatic dispatch at {datetime.fromisoformat(disp['dispatch_at']):%d %b %Y %I:%M %p}"

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="0B2545")
    head_font = Font(bold=True, color="FFFFFF")
    total_fill = PatternFill("solid", fgColor="E8EEF3")
    thin = Side(style="thin", color="CBD5E1")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet_header(ws, title, ncols):
        ws.append([profile.business_name])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"{title} - {d:%A, %d %b %Y}"])
        ws["A2"].font = Font(bold=True, size=12)
        ws.append([state])
        ws.append([dispatch_line])
        ws.append([f"Downloaded {datetime.now():%d %b %Y %I:%M %p}"])
        ws["A5"].font = Font(italic=True, color="64748B")
        for r in range(1, 6):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        ws.append([])

    def table(ws, headers, data, widths, money_cols=(), qty_cols=(), total=None):
        ws.append(headers)
        hr = ws.max_row
        for i, _ in enumerate(headers, 1):
            c = ws.cell(row=hr, column=i)
            c.font, c.fill, c.border = head_font, head_fill, box
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in data:
            ws.append(row)
            for i in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=i).border = box
        if total:
            ws.append(total)
            for i in range(1, len(headers) + 1):
                c = ws.cell(row=ws.max_row, column=i)
                c.font, c.fill, c.border = bold, total_fill, box
        for col in money_cols:
            for r in range(hr + 1, ws.max_row + 1):
                ws.cell(row=r, column=col).number_format = '"Rs "#,##0.00'
        for col in qty_cols:
            for r in range(hr + 1, ws.max_row + 1):
                ws.cell(row=r, column=col).number_format = '#,##0.##'
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=hr + 1, column=1)

    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    # One quantity column per product ordered that day
    prod_cols = [(t["product_id"], f"{t['product_name']} ({t['unit']})") for t in summary["products"]]
    qty_of = lambda r, pid: round(sum(i["quantity"] for i in r["items"] if i["product_id"] == pid), 3) or None
    headers = (["Sr", "Route", "Delivery boy", "Code", "Customer", "Phone", "Address"]
               + [label for _, label in prod_cols] + ["Amount", "Status", "Delivered"])
    first_q = 8
    amount_col = first_q + len(prod_cols)
    sheet_header(ws, "Orders", len(headers))
    data = []
    for n, r in enumerate(rows, 1):
        c = customers.get(r["customer_id"])
        data.append([n, r["route_name"] or "No route", r["driver_name"] or "-", r["customer_code"], r["customer_name"],
                     c.phone if c else "", c.delivery_address if c else ""]
                    + [qty_of(r, pid) for pid, _ in prod_cols]
                    + [r["total_amount"], r["status"].replace("_", " ").title(), r["delivered_what"] or ""])
    total = (["", "", "", "", f"Total ({len(rows)} customers)", "", ""]
             + [round(sum(qty_of(r, pid) or 0 for r in rows), 3) for pid, _ in prod_cols]
             + [round(sum(r["total_amount"] for r in rows), 2), "", ""])
    table(ws, headers, data, [5, 16, 16, 11, 26, 13, 30] + [14] * len(prod_cols) + [13, 16, 30],
          money_cols=(amount_col,), qty_cols=tuple(range(first_q, amount_col)), total=total)

    ws2 = wb.create_sheet("By route")
    headers2 = ["Route", "Delivery boy", "Customers"] + [label for _, label in prod_cols] + ["Amount"]
    sheet_header(ws2, "What to load per route", len(headers2))
    by_route = {}
    for r in rows:
        key = (r["route_name"] or "No route", r["driver_name"] or "-")
        agg = by_route.setdefault(key, {"n": 0, "amount": 0.0, "qty": {}})
        agg["n"] += 1
        agg["amount"] += r["total_amount"]
        for pid, _ in prod_cols:
            agg["qty"][pid] = agg["qty"].get(pid, 0.0) + (qty_of(r, pid) or 0.0)
    data2 = [[k[0], k[1], v["n"]] + [round(v["qty"].get(pid, 0.0), 3) for pid, _ in prod_cols] + [round(v["amount"], 2)]
             for k, v in sorted(by_route.items())]
    table(ws2, headers2, data2, [20, 20, 12] + [14] * len(prod_cols) + [14],
          money_cols=(4 + len(prod_cols),), qty_cols=tuple(range(4, 4 + len(prod_cols))),
          total=["Total", "", len(rows)] + [round(sum(v["qty"].get(pid, 0.0) for v in by_route.values()), 3) for pid, _ in prod_cols]
          + [round(sum(v["amount"] for v in by_route.values()), 2)])

    ws3 = wb.create_sheet("No order")
    sheet_header(ws3, "Active customers with no order", 4)
    data3 = [[customers[x["id"]].customer_code, x["name"], customers[x["id"]].phone or "",
              "Cancelled" if any(c["customer_id"] == x["id"] for c in cancelled) else "No order"]
             for x in summary["not_ordered"]]
    table(ws3, ["Code", "Customer", "Phone", "Status"], data3, [11, 28, 14, 14])

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="Orders_{d:%Y-%m-%d}.xlsx"'},
    )
